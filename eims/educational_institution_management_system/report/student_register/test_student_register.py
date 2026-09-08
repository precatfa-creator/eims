from io import BytesIO
from unittest import TestCase
from unittest.mock import patch

import frappe
from openpyxl import load_workbook

from eims.educational_institution_management_system.report.student_register import student_register as report


class TestStudentRegister(TestCase):
	def test_views_include_expected_details(self):
		self.assertEqual(len(report.get_columns("Directory")), 8)
		self.assertEqual(len(report.get_columns("Full Details")), 25)
		self.assertIn("outstanding_amount", [c["fieldname"] for c in report.get_columns("Fees")])
		self.assertEqual([c["fieldname"] for c in report.get_columns("Roster")], ["student_name"])

	def test_invalid_date_range(self):
		with patch.object(report, "check_access"), self.assertRaises(frappe.ValidationError):
			report.execute({"academic_year": "2025-2026", "from_date": "2026-09-07", "to_date": "2026-01-01"})

	def test_no_permitted_students_returns_no_data(self):
		with patch.object(frappe, "get_list", return_value=[]), patch.object(frappe.db, "sql") as sql:
			self.assertEqual(report.get_data(frappe._dict(academic_year="2025-2026")), [])
			sql.assert_not_called()

	def test_excel_footer_and_formula_safety(self):
		columns = [{"fieldname": "student_name", "label": "Student", "width": 100}]
		with patch.object(report, "execute", return_value=(columns, [{"student_name": "=1+1"}])), patch.object(frappe.permissions, "can_export", return_value=True):
			report.export_xlsx({"academic_year": "2025-2026"})
		ws = load_workbook(BytesIO(frappe.local.response.filecontent)).active
		self.assertEqual(ws['A4'].data_type, 's')
		self.assertEqual(ws['A4'].value, '=1+1')
		self.assertEqual(ws.cell(ws.max_row, 1).value, "Exported on")
		self.assertIn(report.get_system_timezone(), ws.cell(ws.max_row, 2).value)
