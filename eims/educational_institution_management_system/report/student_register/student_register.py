# Copyright (c) 2026, Catfa and Contributors
# See license.txt
#
# The register the office keeps on paper: one row per enrolment for a year, with the
# identity, contact, guardian and fee columns already filled in, so it can be printed
# or exported to Excel without anyone retyping it.

import frappe
import frappe.permissions
from frappe import _
from frappe.utils import getdate, now_datetime, get_system_timezone

# Every filter is a plain equality on one column, so the where clause writes itself.
FILTER_COLUMNS = {
	"student": "s.name",
	"gender": "s.gender",
	"nationality": "s.nationality",
	"stage": "e.stage",
	"section": "e.section",
	"status": "e.status",
	"payment_status": "e.payment_status",
}


def execute(filters=None):
	filters = frappe._dict(filters or {})
	check_access()
	if not filters.academic_year:
		frappe.throw(_("Please select an Academic Year"))
	if filters.from_date and filters.to_date and getdate(filters.from_date) > getdate(filters.to_date):
		frappe.throw(_("From Date must be before To Date"))

	data = get_data(filters)
	return get_columns(filters.get("view")), data, None, get_chart(data), get_report_summary(data)


def get_columns(view=None):
	columns = [
		{"fieldname": "student", "label": "رقم الطالب", "fieldtype": "Link", "options": "EIMS Student", "width": 90},
		{"fieldname": "student_name", "label": "اسم الطالب", "fieldtype": "Data", "width": 220},
		{"fieldname": "student_name_en", "label": "الاسم بالإنجليزية", "fieldtype": "Data", "width": 200},
		{"fieldname": "enrollment_number", "label": "رقم القيد", "fieldtype": "Data", "width": 100},
		{"fieldname": "gender", "label": "الجنس", "fieldtype": "Data", "width": 70},
		{"fieldname": "date_of_birth", "label": "تاريخ الميلاد", "fieldtype": "Date", "width": 110},
		{"fieldname": "place_of_birth", "label": "مكان الميلاد", "fieldtype": "Data", "width": 120},
		{"fieldname": "nationality", "label": "الجنسية", "fieldtype": "Link", "options": "EIMS Nationality", "width": 100},
		{"fieldname": "national_id", "label": "الرقم الوطني", "fieldtype": "Data", "width": 130},
		{"fieldname": "passport_number", "label": "رقم الجواز", "fieldtype": "Data", "width": 110},
		{"fieldname": "stage", "label": "المرحلة", "fieldtype": "Link", "options": "EIMS Stage", "width": 130},
		{"fieldname": "section", "label": "الشعبة", "fieldtype": "Link", "options": "EIMS Section", "width": 120},
		{"fieldname": "enrollment_status", "label": "حالة التسجيل", "fieldtype": "Data", "width": 100},
		{"fieldname": "enrollment_date", "label": "تاريخ التسجيل", "fieldtype": "Date", "width": 110},
		{"fieldname": "phone", "label": "هاتف الطالب", "fieldtype": "Data", "width": 120},
		{"fieldname": "email", "label": "البريد الإلكتروني", "fieldtype": "Data", "width": 180},
		{"fieldname": "address", "label": "العنوان", "fieldtype": "Data", "width": 200},
		{"fieldname": "guardian_name", "label": "ولي الأمر", "fieldtype": "Data", "width": 200},
		{"fieldname": "guardian_phone", "label": "هاتف ولي الأمر", "fieldtype": "Data", "width": 120},
		{"fieldname": "second_guardian_phone", "label": "هاتف بديل", "fieldtype": "Data", "width": 120},
		{"fieldname": "total_fee", "label": "إجمالي الرسوم", "fieldtype": "Currency", "width": 120},
		{"fieldname": "paid_amount", "label": "المدفوع", "fieldtype": "Currency", "width": 120},
		{"fieldname": "outstanding_amount", "label": "المتبقي", "fieldtype": "Currency", "width": 120},
		{"fieldname": "payment_status", "label": "حالة السداد", "fieldtype": "Data", "width": 110},
		{"fieldname": "enrollment", "label": "التسجيل", "fieldtype": "Link", "options": "EIMS Enrollment", "width": 160},
	]

	views = {
		"Directory": {"student", "student_name", "enrollment_number", "stage", "section", "enrollment_status", "guardian_name", "guardian_phone"},
		"Contacts": {"student", "student_name", "section", "phone", "email", "address", "guardian_name", "guardian_phone", "second_guardian_phone"},
		"Fees": {"student", "student_name", "stage", "section", "total_fee", "paid_amount", "outstanding_amount", "payment_status"},
		# The paper roster: names only, the rest of the sheet is filled in by hand.
		"Roster": {"student_name"},
	}
	if view in views:
		return [column for column in columns if column["fieldname"] in views[view]]
	return columns


def check_access():
	if not frappe.get_doc("Report", "Student Register").is_permitted():
		frappe.throw(_("Not permitted"), frappe.PermissionError)
	frappe.has_permission("EIMS Student", "read", throw=True)
	frappe.has_permission("EIMS Enrollment", "read", throw=True)


def get_data(filters):
	conditions = ["e.academic_year = %(academic_year)s", "e.docstatus < 2"]
	conditions += [f"{column} = %({key})s" for key, column in FILTER_COLUMNS.items() if filters.get(key)]

	# Apply document/user permissions before joining personal and financial data.
	students = frappe.get_list("EIMS Student", pluck="name", limit_page_length=0)
	enrollments = frappe.get_list("EIMS Enrollment", filters={"academic_year": filters.academic_year, "docstatus": ["<", 2]}, pluck="name", limit_page_length=0)
	if not students or not enrollments:
		return []
	filters = frappe._dict(filters)
	filters.update({"allowed_students": students, "allowed_enrollments": enrollments})
	conditions += ["s.name in %(allowed_students)s", "e.name in %(allowed_enrollments)s"]
	if filters.from_date:
		conditions.append("e.enrollment_date >= %(from_date)s")
	if filters.to_date:
		conditions.append("e.enrollment_date <= %(to_date)s")

	rows = frappe.db.sql(
		f"""
			select
				s.name as student, s.student_name, s.student_name_en, s.enrollment_number,
				s.gender, s.date_of_birth, s.place_of_birth, s.nationality, s.national_id,
				s.passport_number, s.phone, s.email, s.address,
				coalesce(nullif(s.parent_name, ''), s.guardian_name) as guardian_name,
				coalesce(nullif(s.parent_phone, ''), s.guardian_phone) as guardian_phone,
				s.second_guardian_phone,
				e.name as enrollment, e.stage, e.section, sec.name_ar as section_label,
				e.status as enrollment_status,
				e.enrollment_date, e.total_fee, e.paid_amount, e.outstanding_amount, e.payment_status
			from `tabEIMS Enrollment` e
			inner join `tabEIMS Student` s on s.name = e.student
			left join `tabEIMS Stage` st on st.name = e.stage
			left join `tabEIMS Section` sec on sec.name = e.section
			where {" and ".join(conditions)}
			order by st.`order`, e.section, s.student_name
		""",
		filters,
		as_dict=True,
	)

	for row in rows:
		# "male"/"female" are stored English; ar.csv already carries both.
		row.gender = _(row.gender) if row.gender else None
	return rows


def get_chart(data):
	if not data:
		return None

	per_stage = {}
	for row in data:
		per_stage[row.stage] = per_stage.get(row.stage, 0) + 1

	return {
		"data": {
			"labels": list(per_stage),
			"datasets": [{"name": "عدد الطلبة", "values": list(per_stage.values())}],
		},
		"type": "bar",
	}


def get_report_summary(data):
	if not data:
		return []

	males = sum(1 for row in data if row.gender == _("male"))
	return [
		{"label": "إجمالي الطلبة", "value": len(data), "datatype": "Int"},
		{"label": "ذكور", "value": males, "datatype": "Int"},
		{"label": "إناث", "value": sum(1 for row in data if row.gender == _("female")), "datatype": "Int"},
		{"label": "إجمالي الرسوم", "value": sum(row.total_fee or 0 for row in data), "datatype": "Currency"},
		{"label": "المتبقي", "value": sum(row.outstanding_amount or 0 for row in data), "datatype": "Currency"},
	]


@frappe.whitelist()
def get_branding(academic_year):
	check_access()
	year = frappe.get_doc("EIMS Academic Year", academic_year)
	year.check_permission("read")
	branding = {"institution_name": "", "logo": "/assets/eims/website/school-logo.webp"}
	if year.institution:
		institution = frappe.get_doc("EIMS Institution", year.institution)
		institution.check_permission("read")
		branding.update(institution_name=institution.institution_name, logo=institution.logo or branding["logo"])
	return branding


@frappe.whitelist()
def export_xlsx(filters=None):
	from io import BytesIO
	from openpyxl import Workbook
	from openpyxl.styles import Font, Alignment

	filters = frappe.parse_json(filters) if isinstance(filters, str) else filters
	columns, data = execute(filters)[:2]
	frappe.permissions.can_export("EIMS Student", raise_exception=True)
	frappe.permissions.can_export("EIMS Enrollment", raise_exception=True)
	wb = Workbook()
	ws = wb.active
	ws.title = "Student Register"
	ws.sheet_view.rightToLeft = frappe.local.lang == "ar"
	ws.append([_("Student Register")])
	ws.append([f"{key}: {value}" for key, value in (filters or {}).items() if value])
	ws.append([column["label"] for column in columns])
	for row in data:
		ws.append([row.get(column["fieldname"]) for column in columns])
	ws.freeze_panes = "A4"
	ws.auto_filter.ref = f"A3:{ws.cell(max(3, ws.max_row), len(columns)).coordinate}"
	stamp = now_datetime().strftime("%Y-%m-%d %H:%M:%S")
	ws.append([])
	ws.append([_("Exported on"), f"{stamp} ({get_system_timezone()})"])
	# Treat names/contact details beginning with '=' as text, never formulas.
	for row in ws:
		for cell in row:
			if cell.data_type == "f":
				cell.data_type = "s"
			cell.alignment = Alignment(vertical="top", wrap_text=True)
	for cell in ws[3]:
		cell.font = Font(bold=True)
	for index, column in enumerate(columns, 1):
		ws.column_dimensions[ws.cell(3, index).column_letter].width = max(14, column["width"] / 7)
	output = BytesIO()
	wb.save(output)
	frappe.local.response.filename = f"Student Register {stamp[:10]}.xlsx"
	frappe.local.response.filecontent = output.getvalue()
	frappe.local.response.type = "binary"
